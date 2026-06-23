import json
import os
import tempfile

import frappe
from frappe.tests.utils import FrappeTestCase

from advocacia.advocacia.documentos import (
	_build_context,
	_contexto_acordo,
	_formatar_data_extenso,
	_infer_category,
	_montar_narrativa_pagamento,
	_valor_por_extenso,
	gerar_documentos_em_lote,
	download_generated_document,
	get_document_placeholder_keys,
	get_kits_disponiveis,
	get_placeholders_referencia,
	get_templates_disponiveis,
)
from advocacia.advocacia.tests.test_setup import create_test_acordo, create_test_legal_case
from frappe.utils import getdate


def _ensure_test_escritorio_config(lawyer_name="Advogada Teste"):
	cfg = frappe.get_single("Office Settings")
	cfg.company_name = "Escritorio Teste Advocacia"
	cfg.cnpj = "11222333000181"
	cfg.oab = "OAB/RS 00.000"
	cfg.lawyer_name = lawyer_name
	cfg.lawyer_cpf = "52998224725"
	cfg.lawyer_rg = "1234567890"
	cfg.address = "Rua Teste, 100, Cidade Teste/RS"
	cfg.sia_registration = "00000"
	cfg.bank_name = "Banco Documentos"
	cfg.bank_pix = "52998224725"
	cfg.save(ignore_permissions=True)
	return cfg


class TestDocumentos(FrappeTestCase):
	def tearDown(self):
		frappe.db.rollback()

	def test_get_templates_vazio_ou_lista(self):
		result = get_templates_disponiveis()
		self.assertIsInstance(result, list)

	def test_get_kits_lista(self):
		result = get_kits_disponiveis()
		self.assertIsInstance(result, list)

	def test_get_placeholders_referencia_grupos(self):
		result = get_placeholders_referencia()
		self.assertIsInstance(result, list)
		grupos = [bloco["grupo"] for bloco in result]
		self.assertIn("Escritório", grupos)
		self.assertIn("Cliente", grupos)
		self.assertIn("Processo", grupos)

	def test_valor_por_extenso(self):
		result = _valor_por_extenso(8000)
		self.assertIn("oito mil", result)
		self.assertIn("reais", result)
		self.assertEqual(_valor_por_extenso(0), "")
		self.assertEqual(_valor_por_extenso(None), "")

	def test_contexto_acordo_mixed_due_date_types(self):
		acordo = create_test_acordo(num_parcelas=2)
		acordo.fee_installments[0].due_date = "2026-06-10"
		acordo.fee_installments[1].due_date = getdate("2026-07-10")
		context = _contexto_acordo(acordo)
		self.assertEqual(len(context["acordo_parcelas"]), 2)
		self.assertEqual(context["acordo_parcelas"][0]["due_date"], "2026-06-10")

	def test_narrativa_pagamento_agrupada(self):
		installments = [
			{"payment_condition": "Data fixa", "due_date": "2026-06-10", "amount": 500, "status": "Pendente", "description": "Parcela 1 de 10"},
			{"payment_condition": "Data fixa", "due_date": "2026-07-10", "amount": 500, "status": "Pendente", "description": "Parcela 2 de 10"},
			{"payment_condition": "Data fixa", "due_date": "2026-08-10", "amount": 500, "status": "Pendente", "description": "Parcela 3 de 10"},
			{"payment_condition": "Data fixa", "due_date": "2026-09-10", "amount": 929, "status": "Pendente", "description": "Parcela 4 de 10"},
			{"payment_condition": "Data fixa", "due_date": "2026-10-10", "amount": 929, "status": "Pendente", "description": "Parcela 5 de 10"},
			{"payment_condition": "Data fixa", "due_date": "2026-11-10", "amount": 929, "status": "Pendente", "description": "Parcela 6 de 10"},
			{"payment_condition": "Data fixa", "due_date": "2026-12-10", "amount": 929, "status": "Pendente", "description": "Parcela 7 de 10"},
			{"payment_condition": "Data fixa", "due_date": "2027-01-10", "amount": 929, "status": "Pendente", "description": "Parcela 8 de 10"},
			{"payment_condition": "Data fixa", "due_date": "2027-02-10", "amount": 929, "status": "Pendente", "description": "Parcela 9 de 10"},
			{"payment_condition": "Data fixa", "due_date": "2027-03-10", "amount": 926, "status": "Pendente", "description": "Parcela 10 de 10"},
		]
		result = _montar_narrativa_pagamento(installments)
		self.assertIn("03 (três) parcelas", result)
		self.assertIn("07 (sete) parcelas", result)
		self.assertIn("última parcela", result)
		self.assertIn("dia 10", result)

	def test_narrativa_pagamento_mista(self):
		installments = [
			{"payment_condition": "Data fixa", "due_date": "2026-06-10", "amount": 2000, "status": "Pendente", "description": "Entrada"},
			{"payment_condition": "Na conclusão", "due_date": "", "amount": 6000, "status": "Pendente", "description": "Saldo final"},
		]
		result = _montar_narrativa_pagamento(installments)
		self.assertIn("dois mil reais", result)
		self.assertIn("conclusão do serviço", result)
		self.assertIn("saldo final", result.lower())

	def test_narrativa_pagamento_vazia(self):
		self.assertEqual(_montar_narrativa_pagamento([]), "")
		self.assertEqual(_montar_narrativa_pagamento([{"status": "Cancelado", "amount": 100}]), "")

	def test_build_context_acordo_narrativa(self):
		_ensure_test_escritorio_config()
		servico = create_test_legal_case()
		create_test_acordo(servico=servico.name, total_amount=8000, num_parcelas=2)
		context = _build_context(servico.name)
		self.assertTrue(context["acordo_valor_extenso"])
		self.assertTrue(context["acordo_parcelas"])
		self.assertEqual(len(context["acordo_parcelas"]), 2)

	def test_placeholders_referencia_cobre_contexto(self):
		_ensure_test_escritorio_config()
		servico = create_test_legal_case()
		context = _build_context(servico.name)
		documented = get_document_placeholder_keys()
		missing = sorted(set(context.keys()) - documented)
		self.assertEqual(missing, [], msg=f"Placeholders no contexto sem documentação: {missing}")

	def test_data_extenso_marco_com_cedilha(self):
		self.assertIn("março", _formatar_data_extenso("2026-03-15"))

	def test_build_context_escritorio_e_cliente(self):
		_ensure_test_escritorio_config()
		servico = create_test_legal_case()
		context = _build_context(servico.name)
		self.assertEqual(context["escritorio_advogada"], "Advogada Teste")
		self.assertEqual(context["escritorio_advogada_cpf"], "529.982.247-25")
		self.assertEqual(context["escritorio_advogada_rg"], "1234567890")
		self.assertEqual(context["escritorio_razao_social"], "Escritorio Teste Advocacia")
		self.assertEqual(context["escritorio_cnpj"], "11.222.333/0001-81")
		self.assertEqual(context["escritorio_banco"], "Banco Documentos")
		self.assertEqual(context["escritorio_pix"], "52998224725")
		self.assertTrue(context["cliente_nome"])
		self.assertEqual(context["nome"], context["cliente_nome"])
		self.assertIn("data_hoje", context)
		self.assertIn("data_hoje_extenso", context)

	def test_gerar_documentos_em_lote(self):
		try:
			from docx import Document as DocxDocument
		except ImportError:
			self.skipTest("python-docx não instalado")

		servico = create_test_legal_case()
		template_names = []

		for idx in range(2):
			with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
				doc = DocxDocument()
				doc.add_paragraph(f"Doc {idx}: {{{{ cliente_nome }}}}")
				doc.save(tmp.name)
				tmp_path = tmp.name

			try:
				file_doc = frappe.get_doc(
					{
						"doctype": "File",
						"file_name": f"template_lote_{idx}.docx",
						"file_url": f"/private/files/template_lote_{idx}.docx",
						"is_private": 1,
					}
				)
				with open(tmp_path, "rb") as f:
					file_doc.content = f.read()
				file_doc.save(ignore_permissions=True)

				template = frappe.get_doc(
					{
						"doctype": "Document Template",
						"title": f"Template Lote {idx} {frappe.generate_hash(length=4)}",
						"document_type": "Contrato",
						"template_file": file_doc.file_url,
						"enabled": 1,
					}
				)
				template.insert(ignore_permissions=True)
				template_names.append(template.name)
			finally:
				if os.path.exists(tmp_path):
					os.unlink(tmp_path)

		result = gerar_documentos_em_lote(servico.name, json.dumps(template_names))
		self.assertTrue(result["success"])
		self.assertEqual(result["data"]["total"], 2)
		self.assertEqual(len(result["data"]["gerados"]), 2)
		for item in result["data"]["gerados"]:
			self.assertTrue(item["file_name"].endswith(".docx"))
			self.assertTrue(item["download_key"])

		first = result["data"]["gerados"][0]
		download_generated_document(first["download_key"])
		self.assertEqual(frappe.local.response.type, "download")
		self.assertTrue(first["file_name"].endswith(".docx"))
		self.assertTrue(frappe.local.response.filecontent)

		case_docs = frappe.get_all(
			"Case Document",
			filters={"legal_case": servico.name, "source": "Gerado pelo App"},
		)
		self.assertEqual(len(case_docs), 0)

	def test_infer_category(self):
		cases = (
			({"title": "Procuracao Ad Judicia", "document_type": "Outro", "description": ""}, "Procuração"),
			({"title": "Contrato Honorarios", "document_type": "Contrato", "description": ""}, "Contrato"),
			({"title": "Modelo generico", "document_type": "Recibo", "description": ""}, "Comprovante"),
			({"title": "Doc qualquer", "document_type": "Outro", "description": "peticao inicial"}, "Petição"),
		)
		for fields, expected in cases:
			template = frappe._dict(fields)
			self.assertEqual(_infer_category(template), expected, msg=fields["title"])
